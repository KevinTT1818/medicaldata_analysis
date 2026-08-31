"""Word 导出的回归测试。

导出的价值在于把结果原样交到用户手里，所以既测结构完整，也测「坏输入不该
毁掉整份文档」——一张图取不到，其余内容照常导出。
"""
from __future__ import annotations

import base64
import struct
import unittest
import zlib

from _env import ROOT as _ROOT  # noqa: F401  必须最先导入：它负责设好数据目录

from docx import Document

from app.adapters import heart_failure, mimic_demo, nhanes, uci_heart  # noqa: F401
from app.analyses import describe, regression, registry, survival  # noqa: F401
from app.cdm import etl, schema
from app.report import docx_export
from app.report import runner as report_runner
from app.report import store as report_store
from app.report.models import ReportInput, ReportSection, SectionSnapshot

HF = "heart_failure"
AGE = "person.age_at_index"
SEX = "person.gender"
EF = "measurement.LOINC:10230-1"
DEATH = "outcome.death"


def _png(width: int = 120, height: int = 80) -> bytes:
    """按 PNG 规范拼一张纯色图。手搓十六进制是不可靠的。"""
    def chunk(tag: bytes, data: bytes) -> bytes:
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))

    raw = b"".join(b"\x00" + bytes([70, 180, 187] * width) for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b""))


def _data_url(payload: bytes) -> str:
    return "data:image/png;base64," + base64.b64encode(payload).decode()


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        schema.init()
        etl.import_dataset(HF)
        # 不能叫 cls.run —— 那会盖住 unittest.TestCase.run
        cls.report_run = cls._build_report()

    @classmethod
    def _sections(cls) -> list[ReportSection]:
        return [
            ReportSection(
                title="基线特征", note="按死亡结局分组", dataset=HF,
                analysis="describe.baseline_table",
                params={"variables": [AGE, SEX, EF], "group_by": DEATH},
            ),
            ReportSection(
                title="生存曲线", dataset=HF, analysis="survival.kaplan_meier",
                params={"outcome": DEATH, "group_by": "condition.HF:hypertension"},
            ),
            ReportSection(
                title="多因素 Cox", dataset=HF, analysis="survival.cox",
                params={"outcome": DEATH, "covariates": [AGE, EF]},
            ),
            ReportSection(
                title="变量分布", dataset=HF, analysis="describe.distribution",
                params={"variables": [EF]},
            ),
        ]

    @classmethod
    def _build_report(cls):
        payload = ReportInput(title="导出测试报告", description="四类结果各一节",
                              sections=cls._sections())
        outputs = [report_runner.run_section(s) for s in payload.sections]
        snapshot = [SectionSnapshot(**o["snapshot"]) for o in outputs]
        saved = report_store.save(payload, snapshot=snapshot)
        cls.report_id = saved["id"]
        return report_runner.run_report(saved["id"])

    @staticmethod
    def _read(stream) -> Document:
        return Document(stream)


class StructureTest(Base):
    def test_every_section_becomes_a_heading(self):
        document = self._read(docx_export.build(self.report_run))
        headings = [p.text for p in document.paragraphs
                    if p.style.name.startswith("Heading")]
        self.assertEqual(len(headings), 4)
        self.assertTrue(headings[0].startswith("1. 基线特征"))
        self.assertTrue(headings[2].startswith("3. 多因素 Cox"))

    def test_title_and_description_are_written(self):
        document = self._read(docx_export.build(self.report_run))
        texts = [p.text for p in document.paragraphs]
        self.assertIn("导出测试报告", texts)
        self.assertIn("四类结果各一节", texts)

    def test_each_result_kind_produces_a_table(self):
        document = self._read(docx_export.build(self.report_run))
        headers = [" | ".join(c.text for c in t.rows[0].cells) for t in document.tables]
        joined = "\n".join(headers)
        self.assertIn("变量", joined)          # Table 1
        self.assertIn("中位生存", joined)       # KM 汇总
        self.assertIn("HR", joined)            # Cox 森林图
        self.assertIn("均值 ± SD", joined)      # 分布

    def test_tables_declare_column_widths(self):
        """不设列宽的话 Word 按内容自适应，同一份文档在不同机器上对不齐。"""
        document = self._read(docx_export.build(self.report_run))
        for table in document.tables:
            for cell in table.rows[0].cells:
                self.assertIsNotNone(cell.width)
                self.assertGreater(cell.width.inches, 0)

    def test_fingerprints_are_carried_into_the_document(self):
        """导出的文档要能追回是哪一次分析算出来的。"""
        document = self._read(docx_export.build(self.report_run))
        texts = " ".join(p.text for p in document.paragraphs)
        for section in self.report_run["sections"]:
            fingerprint = section["job"]["spec"]["fingerprint"]
            self.assertIn(fingerprint, texts)

    def test_reproducibility_verdict_is_stated(self):
        document = self._read(docx_export.build(self.report_run))
        texts = " ".join(p.text for p in document.paragraphs)
        self.assertIn("可复现性核对", texts)

    def test_section_notes_are_included(self):
        document = self._read(docx_export.build(self.report_run))
        texts = " ".join(p.text for p in document.paragraphs)
        self.assertIn("按死亡结局分组", texts)


class ImageHandlingTest(Base):
    def test_valid_images_are_embedded(self):
        images = {"1": [_data_url(_png())], "2": [_data_url(_png())]}
        document = self._read(docx_export.build(self.report_run, images))
        self.assertEqual(len(document.inline_shapes), 2)

    def test_multiple_images_per_section(self):
        """一节可能有多张图 —— 森林图带 ROC、分布图每个变量一张。"""
        images = {"3": [_data_url(_png()), _data_url(_png()), _data_url(_png())]}
        document = self._read(docx_export.build(self.report_run, images))
        self.assertEqual(len(document.inline_shapes), 3)

    def test_corrupt_image_does_not_kill_the_document(self):
        """能解出 base64 不代表是有效图片。一张坏图不该毁掉整份文档。"""
        broken = _data_url(b"this decodes fine but is not a png")
        images = {"0": [_data_url(_png())], "1": [broken]}
        document = self._read(docx_export.build(self.report_run, images))

        self.assertEqual(len(document.inline_shapes), 1)     # 好的那张还在
        texts = " ".join(p.text for p in document.paragraphs)
        self.assertIn("无法嵌入", texts)
        # 其余内容照常
        self.assertEqual(
            len([p for p in document.paragraphs if p.style.name.startswith("Heading")]),
            4,
        )

    def test_malformed_data_url_is_skipped_silently(self):
        """连 data URL 都不是的，直接跳过 —— 不该在文档里留噪音。"""
        images = {"0": ["not a data url at all"]}
        document = self._read(docx_export.build(self.report_run, images))
        self.assertEqual(len(document.inline_shapes), 0)
        texts = " ".join(p.text for p in document.paragraphs)
        self.assertNotIn("无法嵌入", texts)

    def test_export_without_any_images_still_works(self):
        document = self._read(docx_export.build(self.report_run, {}))
        self.assertEqual(len(document.inline_shapes), 0)
        self.assertGreater(len(document.tables), 0)


class FailedSectionTest(Base):
    def test_failed_section_is_reported_not_skipped(self):
        """跑不通的一节要在文档里说明，不能悄悄消失 —— 读者会以为报告是完整的。"""
        broken = dict(self.report_run)
        broken["sections"] = list(self.report_run["sections"])
        broken["sections"][1] = {
            "title": "跑不通的一节", "note": None, "status": "failed",
            "error": "DesignError: 完整病例数太少", "job": {"spec": {}}, "result": None,
        }
        document = self._read(docx_export.build(broken))
        texts = " ".join(p.text for p in document.paragraphs)
        self.assertIn("跑不通的一节", texts)
        self.assertIn("完整病例数太少", texts)


if __name__ == "__main__":
    unittest.main()
