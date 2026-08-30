"""把报告导出成 Word 文档。

图表由前端把 ECharts 画布转成 PNG 一起送上来 —— 服务端不重画。
重画意味着引入一套绘图依赖，还要让它画得和用户屏幕上看到的一模一样，
那是两份必然会互相偏离的实现。用户看到什么，文档里就是什么。

表格全部显式设列宽（DXA）：不设的话 Word 会按内容自适应，
同一份文档在不同机器上列宽不同，对不齐。
"""
from __future__ import annotations

import base64
import binascii
import io
import re
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

#: 正文可用宽度（Letter 纵向，左右各 1 英寸页边距）
CONTENT_WIDTH = Inches(6.5)

MUTED = RGBColor(0x5C, 0x6E, 0x75)
ACCENT = RGBColor(0x0A, 0x5B, 0x63)
WARN = RGBColor(0x8F, 0x5D, 0x10)

_DATA_URL = re.compile(r"^data:image/(png|jpeg|jpg);base64,(.+)$", re.DOTALL)


def _decode_image(data_url: str) -> io.BytesIO | None:
    """解析前端送来的 data URL。解不开就跳过这张图，不让整份导出失败。"""
    match = _DATA_URL.match(data_url.strip())
    if not match:
        return None
    try:
        return io.BytesIO(base64.b64decode(match.group(2)))
    except (binascii.Error, ValueError):
        return None


def _muted(paragraph, text: str, size: int = 8, color: RGBColor = MUTED) -> None:
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color


def _table(document, headers: list[str], rows: list[list[str]],
           widths: list[float] | None = None):
    """建一张表。列宽必须显式给，否则 Word 按内容自适应，各处对不齐。"""
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    if widths is None:
        widths = [1.0 / len(headers)] * len(headers)
    total = CONTENT_WIDTH.inches
    column_widths = [Inches(total * w) for w in widths]

    for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(9)
        cell.width = column_widths[index]

    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(str(value))
            run.font.size = Pt(9)
            cells[index].width = column_widths[index]

    return table


def _format_p(value: float | None) -> str:
    if value is None:
        return "—"
    return "<0.001" if value < 0.001 else f"{value:.3f}"


# ---------------------------------------------------------------- 各类结果

def _write_baseline_table(document, payload: dict[str, Any]) -> None:
    groups = payload.get("groups") or []
    show_tests = bool(groups) and not payload.get("weighted")
    has_fmi = any(r.get("fmi") is not None for r in payload["rows"])

    headers = ["变量", f"总体 (n={payload['overall_n']})"]
    headers += [f"{g['label']} (n={g['n']})" for g in groups]
    headers.append("缺失")
    if show_tests or payload.get("weighted"):
        headers += ["p", "p 校正"]
        if has_fmi:
            headers.append("FMI")
        headers.append("检验方法")

    keys = ["__overall__"] + [g["key"] for g in groups]
    rows: list[list[str]] = []

    for row in payload["rows"]:
        unit = f"（{row['unit']}）" if row.get("unit") else ""
        stat = {"mean_sd": "均值 ± SD", "median_iqr": "中位数 (IQR)"}.get(
            row["stat"], "n (%)")
        line = [f"{row['label']}{unit}，{stat}"]
        line += [row["cells"].get(key, "") if row["kind"] == "continuous" else ""
                 for key in keys]
        line.append(str(row["n_missing"]))
        if show_tests or payload.get("weighted"):
            line += [_format_p(row["p"]), _format_p(row["p_adj"])]
            if has_fmi:
                line.append("—" if row.get("fmi") is None else f"{row['fmi']:.3f}")
            line.append(row.get("test") or "—")
        rows.append(line)

        for level in row.get("levels") or []:
            sub = [f"    {level['label']}"]
            sub += [level["cells"].get(key, "") for key in keys]
            sub += [""] * (len(headers) - len(sub))
            rows.append(sub)

    first = 0.30
    rest = (1 - first) / max(len(headers) - 1, 1)
    _table(document, headers, rows, [first] + [rest] * (len(headers) - 1))


def _write_forest(document, payload: dict[str, Any]) -> None:
    effect = payload.get("effect_label", "效应量")
    level = int(payload.get("conf_level", 0.95) * 100)
    has_fmi = any(r.get("fmi") is not None for r in payload["rows"])
    has_ph = bool(payload.get("ph_test"))
    ph_by_label = {t["label"]: t["p"] for t in payload.get("ph_test") or []}

    headers = ["项目", effect, f"{level}% CI", "p"]
    if has_ph:
        headers.append("PH p")
    if has_fmi:
        headers.append("FMI")

    rows = []
    for row in payload["rows"]:
        label = row["label"]
        if row.get("reference"):
            label += f"（参照：{row['reference']}）"
        line = [
            label,
            f"{row['estimate']:.3f}",
            f"{row['ci_lower']:.3f}–{row['ci_upper']:.3f}",
            _format_p(row["p"]),
        ]
        if has_ph:
            line.append(_format_p(ph_by_label.get(row["label"])))
        if has_fmi:
            line.append("—" if row.get("fmi") is None else f"{row['fmi']:.3f}")
        rows.append(line)

    first = 0.34
    rest = (1 - first) / max(len(headers) - 1, 1)
    _table(document, headers, rows, [first] + [rest] * (len(headers) - 1))

    stats_line = [f"纳入 {payload['n_used']} 例", f"事件 {payload['n_events']} 例"]
    if payload.get("n_dropped"):
        stats_line.append(f"因缺失剔除 {payload['n_dropped']} 例")
    if payload.get("concordance") is not None:
        stats_line.append(f"C-index {payload['concordance']:.3f}")
    if payload.get("roc", {}).get("auc") is not None:
        stats_line.append(f"AUC {payload['roc']['auc']:.3f}")
    _muted(document.add_paragraph(), " · ".join(stats_line))


def _write_km(document, payload: dict[str, Any]) -> None:
    unit = payload.get("time_unit", "天")
    rows = [
        [
            s["name"], str(s["n"]), str(s["events"]), str(s["censored"]),
            "未达到" if s["median_survival"] is None
            else f"{s['median_survival']} {unit}",
        ]
        for s in payload["series"]
    ]
    _table(document, ["分组", "n", "事件", "删失", "中位生存"], rows,
           [0.32, 0.17, 0.17, 0.17, 0.17])

    if payload.get("logrank_p") is not None:
        paragraph = document.add_paragraph()
        _muted(paragraph, f"{payload['test']}　p = {_format_p(payload['logrank_p'])}",
               size=9, color=ACCENT)

    at_risk = payload.get("at_risk") or {}
    if at_risk.get("times"):
        document.add_paragraph()
        _muted(document.add_paragraph(), f"各时点尚在随访人数（{unit}）")
        headers = ["分组"] + [str(t) for t in at_risk["times"]]
        rows = [[r["name"]] + [str(c) for c in r["counts"]] for r in at_risk["rows"]]
        first = 0.24
        rest = (1 - first) / max(len(headers) - 1, 1)
        _table(document, headers, rows, [first] + [rest] * (len(headers) - 1))


def _write_distribution(document, payload: dict[str, Any]) -> None:
    for panel in payload["panels"]:
        unit = f"（{panel['unit']}）" if panel.get("unit") else ""
        _muted(document.add_paragraph(), f"{panel['label']}{unit}", size=9)
        if panel["kind"] == "continuous":
            rows = []
            for series in panel["series"]:
                box = series.get("box")
                if not box:
                    continue
                sd = "—" if box["sd"] is None else f"{box['sd']:.2f}"
                rows.append([
                    series["name"], str(box["n"]),
                    f"{box['mean']:.2f} ± {sd}",
                    f"{box['median']:.2f}",
                    f"{box['q1']:.2f}–{box['q3']:.2f}",
                    str(len(box["outliers"])),
                ])
            _table(document, ["分组", "n", "均值 ± SD", "中位数", "IQR", "离群"],
                   rows, [0.22, 0.12, 0.24, 0.14, 0.19, 0.09])
        else:
            headers = ["分组"] + list(panel.get("levels") or [])
            rows = [[s["name"]] + [str(c) for c in s["counts"]]
                    for s in panel["series"]]
            first = 0.24
            rest = (1 - first) / max(len(headers) - 1, 1)
            _table(document, headers, rows, [first] + [rest] * (len(headers) - 1))


def _write_missingness(document, payload: dict[str, Any]) -> None:
    _muted(document.add_paragraph(),
           f"完整记录 {payload['complete_cases']} / {payload['overall_n']}"
           f"（{payload['complete_pct']}%）", size=9)
    rows = [[r["label"], str(r["n_missing"]), f"{r['pct_missing']}%"]
            for r in payload["rows"]]
    _table(document, ["变量", "缺失数", "缺失率"], rows, [0.5, 0.25, 0.25])


_WRITERS = {
    "baseline_table": _write_baseline_table,
    "forest": _write_forest,
    "km_curve": _write_km,
    "distribution": _write_distribution,
    "missingness": _write_missingness,
}


# ---------------------------------------------------------------- 入口

def build(
    run: dict[str, Any], images: dict[str, list[str]] | None = None
) -> io.BytesIO:
    """把一次报告运行结果渲染成 .docx，返回内存中的字节流。

    images 的键是小节序号，值是该节的图片列表 —— 一节可能有多张图
    （森林图带 ROC、分布图每个变量一张）。
    """
    images = images or {}
    document = Document()

    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)      # Letter，python-docx 默认是 Letter，显式写明
    section.page_height = Inches(11)
    for attribute in ("left_margin", "right_margin"):
        setattr(section, attribute, Inches(1))

    meta = run["report"]
    document.add_heading(meta["title"], level=0)
    if meta.get("description"):
        paragraph = document.add_paragraph(meta["description"])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    _muted(document.add_paragraph(),
           f"保存于 {meta['created_at'][:19]}　更新于 {meta['updated_at'][:19]}")

    comparisons = {c["index"]: c for c in run.get("comparisons", [])}
    if run.get("has_baseline"):
        verdict = "与保存时完全一致" if run["reproducible"] else "与保存时相比有变化"
        paragraph = document.add_paragraph()
        _muted(paragraph, f"可复现性核对：{verdict}", size=9,
               color=ACCENT if run["reproducible"] else WARN)

    for index, section_result in enumerate(run["sections"]):
        document.add_heading(f"{index + 1}. {section_result['title']}", level=1)

        comparison = comparisons.get(index)
        if comparison:
            label = {"match": "与基线一致", "changed": "数字有变",
                     "failed": "跑不通", "no_baseline": "无基线"}.get(
                comparison["verdict"], comparison["verdict"])
            _muted(document.add_paragraph(),
                   f"{label}　{comparison['detail']}", size=8,
                   color=WARN if comparison["verdict"] == "changed" else MUTED)

        if section_result["status"] == "failed":
            _muted(document.add_paragraph(),
                   f"这一节没有跑通：{section_result.get('error')}", size=9, color=WARN)
            continue

        for data_url in images.get(str(index), []):
            stream = _decode_image(data_url)
            if stream is None:
                continue
            try:
                document.add_picture(stream, width=CONTENT_WIDTH)
            except Exception:  # noqa: BLE001
                # 能解出 base64 不代表是有效图片。一张坏图不该毁掉整份文档，
                # 跳过并在原位说明，其余内容照常导出。
                _muted(document.add_paragraph(),
                       "（这一节的图表无法嵌入，请回到网页查看）", size=8, color=WARN)

        result = section_result.get("result") or {}
        writer = _WRITERS.get(result.get("kind"))
        if writer:
            writer(document, result)

        if section_result.get("note"):
            paragraph = document.add_paragraph()
            _muted(paragraph, section_result["note"], size=9, color=MUTED)

        for note in result.get("notes") or []:
            _muted(document.add_paragraph(), f"· {note}", size=8)

        job = section_result.get("job") or {}
        fingerprint = (job.get("spec") or {}).get("fingerprint")
        if fingerprint:
            detail = f"指纹 {fingerprint}"
            if job.get("cohort_n") is not None:
                detail += f"　{job['cohort_n']} / {job['cohort_total']} 例"
            _muted(document.add_paragraph(), detail, size=8)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
